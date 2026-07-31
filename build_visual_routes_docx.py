from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Administrator\Desktop\vetai-complete2")
OUT = ROOT / "犬口腔AI研究两条路线_直观版.docx"
ASSET = ROOT / "docx_visual_assets"
ASSET.mkdir(exist_ok=True)

FONT = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"


def fnt(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def rounded(draw, xy, fill, outline=None, radius=24, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, box, text, font, fill="#17324D", spacing=8):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2), text, font=font, fill=fill,
                        spacing=spacing, align="center")


def arrow(draw, x1, y1, x2, y2, color="#6B7A89", width=6):
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    if abs(y2 - y1) >= abs(x2 - x1):
        draw.polygon([(x2, y2), (x2 - 13, y2 - 20), (x2 + 13, y2 - 20)], fill=color)
    else:
        draw.polygon([(x2, y2), (x2 - 20, y2 - 13), (x2 - 20, y2 + 13)], fill=color)


def make_overview(path):
    im = Image.new("RGB", (1600, 760), "white")
    d = ImageDraw.Draw(im)
    d.text((80, 35), "我建议的研究顺序", font=fnt(46, True), fill="#17324D")
    d.text((80, 98), "先回答“哪个模型更好”，再回答“它能否达到兽医水平”", font=fnt(27), fill="#5B6770")

    boxes = [
        (90, 220, 500, 560, "阶段 1｜路线B", "四模型公平比较", "200张图 + 统一Prompt\n筛出综合最佳1–2个模型", "#EAF4FF", "#2D78C4"),
        (600, 220, 1010, 560, "筛选", "保留最佳模型", "不是只看总准确率\n同时看稳定性、成本和分体征表现", "#FFF7E6", "#D79A27"),
        (1110, 220, 1520, 560, "阶段 2｜路线A", "AI与兽医比较", "专家金标准 + 多名兽医\n检验AI是否达到临床辅助水平", "#EBF8F1", "#2D9461"),
    ]
    for x1, y1, x2, y2, tag, title, detail, fill, line in boxes:
        rounded(d, (x1, y1, x2, y2), fill, line, 30, 4)
        d.text((x1 + 32, y1 + 30), tag, font=fnt(24, True), fill=line)
        center_text(d, (x1 + 25, y1 + 95, x2 - 25, y1 + 190), title, fnt(38, True))
        center_text(d, (x1 + 35, y1 + 195, x2 - 35, y2 - 25), detail, fnt(24), fill="#425466", spacing=13)
    arrow(d, 515, 390, 580, 390)
    arrow(d, 1025, 390, 1090, 390)
    rounded(d, (260, 635, 1340, 715), "#17324D", radius=20)
    center_text(d, (260, 635, 1340, 715), "当前只有照片和API：优先路线B　｜　有专家、兽医和临床资料：再做路线A", fnt(25, True), fill="white")
    im.save(path, quality=95)


def make_route(path, route="A"):
    is_a = route == "A"
    accent = "#2D9461" if is_a else "#2D78C4"
    pale = "#EBF8F1" if is_a else "#EAF4FF"
    title = "路线A｜AI与兽医临床能力比较" if is_a else "路线B｜四个模型严格横向比较"
    question = "AI是否达到初级/临床兽医的体征识别水平？" if is_a else "相同输入条件下，哪个模型更准、更稳、更划算？"
    steps = (
        [
            ("1", "界定用途", "筛查或第二意见，不主张自主诊断"),
            ("2", "建立金标准", "≥3名专家；尽量结合检查、X线或病理"),
            ("3", "设置人类组", "3–5名初级兽医；可加普通临床兽医"),
            ("4", "统一信息", "AI和兽医看到完全相同的照片与病史"),
            ("5", "独立盲法判读", "统一9项体征、三分类标签，不互看答案"),
            ("6", "预设主要终点", "推荐宏平均平衡准确率；其余为次要指标"),
            ("7", "直接比较差值", "计算AI－兽医差值及95%CI"),
            ("8", "判断临床价值", "采用非劣效设计，说明可用与不可用体征"),
        ] if is_a else [
            ("1", "确认病例单位", "按犬只/病例计数；同犬多图视为聚类"),
            ("2", "开发集锁Prompt", "总计200张时，可用20–30张开发、其余测试"),
            ("3", "锁定模型版本", "记录真实API ID、日期、推理和图像设置"),
            ("4", "保持输入一致", "相同图片字节、Prompt和JSON格式"),
            ("5", "一病例一请求", "禁用记忆、联网和连续对话污染"),
            ("6", "主运行一次", "所有模型对全部锁定测试病例运行"),
            ("7", "稳定性复测", "分层抽40张，每模型额外运行2次"),
            ("8", "正式配对统计", "比较总体差异及 model × clinical sign"),
            ("9", "综合评价", "准确性、稳定性、成本、速度和错误类型"),
        ]
    )
    h = 1220 if is_a else 1330
    im = Image.new("RGB", (1600, h), "white")
    d = ImageDraw.Draw(im)
    d.text((75, 35), title, font=fnt(42, True), fill="#17324D")
    rounded(d, (75, 105, 1525, 205), pale, accent, 20, 3)
    d.text((110, 135), "核心问题", font=fnt(25, True), fill=accent)
    d.text((310, 133), question, font=fnt(28, True), fill="#17324D")
    y = 245
    for i, (num, name, detail) in enumerate(steps):
        rounded(d, (120, y, 1480, y + 92), "#F8FAFC", "#D9E2EA", 18, 2)
        d.ellipse((148, y + 18, 204, y + 74), fill=accent)
        center_text(d, (148, y + 18, 204, y + 74), num, fnt(24, True), fill="white")
        d.text((245, y + 18), name, font=fnt(25, True), fill="#17324D")
        d.text((510, y + 20), detail, font=fnt(23), fill="#4D5C68")
        if i < len(steps) - 1:
            arrow(d, 176, y + 94, 176, y + 116, color="#A6B4C0", width=4)
        y += 108
    note = "关键：直接检验AI与兽医的性能差，不能比较两组各自的P值。" if is_a else "关键：测试集开始后不再改Prompt；结论不能只写排行榜。"
    rounded(d, (120, y + 15, 1480, y + 100), accent, radius=20)
    center_text(d, (120, y + 15, 1480, y + 100), note, fnt(25, True), fill="white")
    im.save(path, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run(run, size=11, bold=False, color="243746"):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    for name, size, color, before, after in [
        ("Title", 25, "17324D", 0, 7),
        ("Heading 1", 17, "2D78C4", 15, 7),
        ("Heading 2", 13, "17324D", 10, 5),
    ]:
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)


def add_page_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(p.add_run(title), 25, True, "17324D")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_run(p.add_run(subtitle), 11.5, False, "5B6770")


def add_image(doc, path, width=7.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(path), width=Inches(width))


def build_doc():
    overview = ASSET / "overview.png"
    route_a = ASSET / "route_a.png"
    route_b = ASSET / "route_b.png"
    make_overview(overview)
    make_route(route_a, "A")
    make_route(route_b, "B")

    doc = Document()
    style_doc(doc)
    sec = doc.sections[0]
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run("犬口腔多模态AI研究｜研究设计速览"), 8.5, False, "7B8790")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("用于研究设计讨论｜2026"), 8, False, "8B949B")

    add_page_title(doc, "犬口腔AI研究：两条路线怎么走", "面向约200张口腔照片与4个多模态模型的直观研究方案")
    add_image(doc, overview, 7.0)

    doc.add_heading("两条路线，一眼看懂", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.25, 2.85, 2.9]
    headers = ["比较项", "路线A：AI vs 兽医", "路线B：模型 vs 模型"]
    for j, cell in enumerate(table.rows[0].cells):
        cell.width = Inches(widths[j])
        set_cell_shading(cell, "17324D")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(headers[j]), 9.5, True, "FFFFFF")
    rows = [
        ("回答什么", "AI能否达到兽医水平", "四个模型谁更好"),
        ("需要资源", "专家 + 多名兽医 + 临床金标准", "照片 + 四个平台API"),
        ("主要统计", "非劣效/直接性能差", "配对比较 + 模型×体征"),
        ("当前建议", "作为第二阶段", "优先实施"),
    ]
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].width = Inches(widths[j])
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 0:
                set_cell_shading(cells[j], "F3F7FA")
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(value), 9.5, j == 0, "243746")

    doc.add_page_break()
    add_page_title(doc, "路线A｜AI和兽医临床能力比较", "重点是临床可用性：AI是否可以成为筛查或第二意见工具")
    add_image(doc, route_a, 6.85)

    doc.add_heading("最终应得出的结论", level=1)
    p = doc.add_paragraph()
    set_run(p.add_run("不是“AI排名第几”，而是："), 11, True, "17324D")
    set_run(p.add_run("AI相对初级/临床兽医的性能差是多少，95%置信区间是否支持非劣效，以及哪些体征仍不安全。"), 11, False, "243746")

    doc.add_page_break()
    add_page_title(doc, "路线B｜四个模型严格横向比较", "重点是公平benchmark：同图、同Prompt、同输出格式、同统计规则")
    add_image(doc, route_b, 6.12)

    doc.add_heading("200张照片怎么分", level=1)
    p = doc.add_paragraph()
    set_run(p.add_run("如果总共只有200张："), 11, True, "17324D")
    set_run(p.add_run("20–30张开发集用于锁定Prompt，170–180张作为正式测试集。若能另找开发图片，最好把200张全部留作锁定测试。"), 11, False, "243746")

    doc.add_heading("最重要的执行规则", level=1)
    for text in [
        "同一只犬的多张照片必须保留 dog_id，统计时按犬只聚类。",
        "每个病例独立请求；四个模型使用完全相同的图片文件和体征定义。",
        "测试开始后不再修改Prompt；所有失败、拒答和 indeterminate 都保留。",
        "除准确率外，同时报告稳定性、成本、速度、假阴性和分体征表现。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        set_run(p.add_run(text), 9.8, False, "243746")

    doc.add_page_break()
    add_page_title(doc, "我的落地建议", "把复杂研究拆成两个连续、可以分别投稿的阶段")
    add_image(doc, overview, 6.9)
    doc.add_heading("现在先做什么", level=1)
    steps = [
        ("01", "整理病例表", "给每张图建立 dog_id、image_id、图像质量和9项专家标签。"),
        ("02", "建立开发集", "用20–30张锁定Prompt、输出JSON和图片预处理规则。"),
        ("03", "跑四模型测试", "完成路线B，筛出综合表现最好的1–2个模型。"),
        ("04", "组织人类评价", "有专家和兽医资源后，使用最佳模型进入路线A。"),
    ]
    for num, name, detail in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        set_run(p.add_run(num + "  "), 13, True, "2D78C4")
        set_run(p.add_run(name + "｜"), 11, True, "17324D")
        set_run(p.add_run(detail), 10.5, False, "243746")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("一句话结论：当前资源先走路线B；路线A留给有兽医团队和可靠金标准之后。"), 12, True, "2D9461")

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
