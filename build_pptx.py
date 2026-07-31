# -*- coding: utf-8 -*-
"""把两张 SVG 信息图原样重建为可编辑 PPT（python-pptx）。"""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

PX = 6477   # EMU per svg unit (1200px -> 8.5in)
FS = 0.51   # pt per svg font unit

WHITE='FFFFFF'; GREEN='0F6E56'; AMBER='B8861E'; INK='1A2B25'; INK2='5C6B65'
INK3='8A9690'; PANEL='F7F9F8'; BORDER='E3EAE6'; DIVIDER='EDEFE9'; AMBER_BG='FBF6E8'

def E(v): return int(v * PX)
def C(h): return RGBColor.from_string(h)
def _noshadow(s): s.shadow.inherit = False

def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = C(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = C(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    _noshadow(s); return s

def rrect(slide, x, y, w, h, fill=None, line=None, lw=1.0, adj=0.08):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    try: s.adjustments[0] = adj
    except Exception: pass
    if fill: s.fill.solid(); s.fill.fore_color.rgb = C(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = C(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    _noshadow(s); return s

def _setrun(r, txt, size, color, bold):
    r.text = txt
    r.font.size = Pt(round(size*FS, 1))
    r.font.bold = bold
    r.font.color.rgb = C(color)
    r.font.name = 'PingFang SC'

def txt_l(slide, x, y, t, size, color, bold=False, w=None):
    if w is None: w = 1140 - x
    tb = slide.shapes.add_textbox(E(x), E(y - size*0.82), E(w), E(size*1.3))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    _setrun(p.add_run(), t, size, color, bold); return tb

def txt_cx(slide, cx, y, t, size, color, bold=False, w=600):
    tb = slide.shapes.add_textbox(E(cx - w/2), E(y - size*0.82), E(w), E(size*1.3))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _setrun(p.add_run(), t, size, color, bold); return tb

def txt_cc(slide, cx, cy, t, size, color, bold=False, w=600):
    tb = slide.shapes.add_textbox(E(cx - w/2), E(cy - size*0.62), E(w), E(size*1.25))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _setrun(p.add_run(), t, size, color, bold); return tb

def circle(slide, cx, cy, r, fill=None, line=None, lw=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(cx-r), E(cy-r), E(r*2), E(r*2))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = C(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = C(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    _noshadow(s); return s

def line(slide, x1, y1, x2, y2, color, lw=1.0, arrow=False, dash=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    cn.line.color.rgb = C(color); cn.line.width = Pt(lw)
    ln = cn.line._get_or_add_ln()
    if arrow:
        ln.append(ln.makeelement(qn('a:tailEnd'), {'type':'triangle','w':'med','len':'med'}))
    if dash:
        ln.append(ln.makeelement(qn('a:prstDash'), {'val':'dash'}))
    _noshadow(cn); return cn

def build_page1(slide):
    rect(slide, 0, 0, 1200, 1820, fill=WHITE)
    rect(slide, 0, 0, 1200, 6, fill=GREEN)
    txt_l(slide, 60, 92, 'VetAI 产品使用逻辑', 50, INK, bold=True)
    txt_l(slide, 60, 138, '对话建档 → 智能诊断 → 报告追问 · 全流程', 28, INK2)
    line(slide, 60, 168, 1140, 168, DIVIDER)
    txt_l(slide, 60, 208, '入口', 28, INK3, bold=True)
    rrect(slide, 60, 226, 520, 92, fill=WHITE, line=BORDER, lw=1.0, adj=0.1)
    rect(slide, 60, 226, 4, 92, fill=GREEN)
    circle(slide, 112, 272, 22, fill=WHITE, line=GREEN, lw=1.4)
    txt_cc(slide, 112, 272, '新', 28, GREEN, bold=True, w=60)
    txt_l(slide, 150, 262, '新宠物', 30, INK, bold=True)
    txt_l(slide, 150, 298, '从问候建档开始，一次性收集档案', 28, INK2)
    rrect(slide, 620, 226, 520, 92, fill=WHITE, line=BORDER, lw=1.0, adj=0.1)
    rect(slide, 620, 226, 4, 92, fill=AMBER)
    circle(slide, 672, 272, 22, fill=WHITE, line=AMBER, lw=1.4)
    txt_cc(slide, 672, 272, '档', 28, AMBER, bold=True, w=60)
    txt_l(slide, 710, 262, '已有档案', 30, INK, bold=True)
    txt_l(slide, 710, 298, '跳过问候建档，直接进入主诉', 28, INK2)
    line(slide, 320, 318, 320, 348, INK3, lw=1.6)
    line(slide, 880, 318, 880, 348, INK3, lw=1.6)
    line(slide, 320, 348, 600, 348, INK3, lw=1.6)
    line(slide, 600, 348, 600, 372, INK3, lw=1.6, arrow=True)
    txt_l(slide, 60, 402, '对话建档阶段流', 28, INK3, bold=True)
    stages = [
        (420,'1','问候建档','收集名字、种类、性别、年龄、品种、体重','缺字段时用快捷选项补问，不走模型',GREEN,'轻量模型 · 2 次调用','提取信息（低随机性，输出结构化）','生成下一句提问',GREEN,'gray'),
        (614,'2','主诉采集','提取主诉、患处部位、症状、持续时间','部位和症状从固定选项中识别',GREEN,'轻量模型 · 2 次调用','提取主诉与症状','生成下一句提问',GREEN,'gray'),
        (808,'3','患处拍照','上传患处照片，落地为本地路径','可选病灶圈选，辅助定位',INK3,'本地处理 · 无调用','图片质量门控','病灶圈选（可选）',INK3,'amber'),
        (1002,'4','分析诊断','进入模型诊断链路','两阶段路由，1~2 次模型调用',AMBER,'见模型调用逻辑','轻量模型预检 + 主力模型诊断','知识库命中时只调 1 次',AMBER,'gray',True),
        (1196,'5','结果输出','格式化诊断报告','含免责声明，非确诊',INK3,'本地处理 · 无调用','结论 + 护理建议','就医指征',INK3,'gray'),
        (1390,'6','追问对话','基于诊断结果回答用户追问','超范围问题建议就医',GREEN,'轻量模型 · 1 次调用','生成回答（不超过 3 句）','语气自然亲切',GREEN,None),
    ]
    for st in stages:
        y, num, name, d1, d2, bar, at, a2, a3, abar, arr, *rest = st
        special = len(rest) > 0 and rest[0]
        cf = AMBER_BG if special else WHITE
        cl = AMBER if special else BORDER
        cw = 1.4 if special else 1.0
        rrect(slide, 60, y, 1080, 168, fill=cf, line=cl, lw=cw, adj=0.07)
        rect(slide, 60, y, 4, 168, fill=bar)
        circle(slide, 112, y+84, 32, fill=(bar if special else WHITE), line=bar, lw=1.6)
        txt_cc(slide, 112, y+84, num, 30, (WHITE if special else bar), bold=True, w=64)
        txt_l(slide, 168, y+58, name, 32, INK, bold=True)
        txt_l(slide, 168, y+94, d1, 28, INK2)
        txt_l(slide, 168, y+130, d2, 28, INK3)
        af = WHITE if special else PANEL
        al = AMBER if special else BORDER
        aw = 1.2 if special else 1.0
        rrect(slide, 700, y+26, 410, 118, fill=af, line=al, lw=aw, adj=0.07)
        rect(slide, 700, y+26, 3, 118, fill=abar)
        txt_l(slide, 720, y+56, at, 28, abar, bold=True)
        txt_l(slide, 720, y+90, a2, 28, INK)
        txt_l(slide, 720, y+124, a3, 28, INK2)
        if arr == 'gray':
            line(slide, 600, y+168, 600, y+194, INK3, lw=1.6, arrow=True)
        elif arr == 'amber':
            line(slide, 600, y+168, 600, y+194, AMBER, lw=2.0, arrow=True)
    line(slide, 60, 1600, 1140, 1600, DIVIDER)
    txt_l(slide, 60, 1640, '核心产出', 28, INK3, bold=True)
    rrect(slide, 60, 1658, 340, 92, fill=WHITE, line=GREEN, lw=1.0, adj=0.1)
    txt_cx(slide, 230, 1690, '宠物档案', 30, GREEN, bold=True, w=300)
    txt_cx(slide, 230, 1726, '可保存复用', 28, INK2, w=300)
    rrect(slide, 430, 1658, 340, 92, fill=WHITE, line=AMBER, lw=1.0, adj=0.1)
    txt_cx(slide, 600, 1690, '诊断报告', 30, AMBER, bold=True, w=300)
    txt_cx(slide, 600, 1726, '结论+护理+就医', 28, INK2, w=300)
    rrect(slide, 800, 1658, 340, 92, fill=WHITE, line=BORDER, lw=1.0, adj=0.1)
    txt_cx(slide, 970, 1690, '追问记录', 30, INK, bold=True, w=300)
    txt_cx(slide, 970, 1726, '对话留痕', 28, INK2, w=300)
    rrect(slide, 60, 1800, 14, 14, fill=GREEN, adj=0.2)
    txt_l(slide, 82, 1815, '模型调用', 22, INK3)
    rrect(slide, 210, 1800, 14, 14, fill=INK3, adj=0.2)
    txt_l(slide, 232, 1815, '本地处理', 22, INK3)
    rrect(slide, 350, 1800, 14, 14, fill=AMBER, adj=0.2)
    txt_l(slide, 372, 1815, '关键节点', 22, INK3)

def build_page2(slide):
    rect(slide, 0, 0, 1200, 1860, fill=WHITE)
    rect(slide, 0, 0, 1200, 6, fill=GREEN)
    txt_l(slide, 60, 92, '模型调用逻辑', 50, INK, bold=True)
    txt_l(slide, 60, 138, '两阶段路由 · 云函数代理 · 模型清单', 28, INK2)
    line(slide, 60, 168, 1140, 168, DIVIDER)
    txt_l(slide, 60, 208, '诊断链路', 28, INK3, bold=True)
    rrect(slide, 320, 226, 560, 76, fill=WHITE, line=BORDER, lw=1.0, adj=0.1)
    txt_cx(slide, 600, 262, '诊断入口', 32, INK, bold=True, w=520)
    txt_cx(slide, 600, 292, '默认开启知识库与快路径', 28, INK2, w=520)
    line(slide, 600, 302, 600, 328, INK3, lw=1.6, arrow=True)
    rrect(slide, 280, 328, 640, 140, fill=PANEL, line=GREEN, lw=1.4, adj=0.07)
    rect(slide, 280, 328, 4, 140, fill=GREEN)
    rrect(slide, 310, 350, 118, 40, fill=GREEN, adj=0.5)
    txt_cc(slide, 369, 371, '调用 1', 28, WHITE, bold=True, w=118)
    txt_l(slide, 448, 372, '视觉特征提取', 32, INK, bold=True)
    txt_l(slide, 310, 420, 'qwen3.6-flash · 预检', 28, GREEN)
    txt_l(slide, 310, 452, '提取视觉特征 · 超时 18 秒 · 轻量快速', 28, INK2)
    line(slide, 600, 468, 600, 494, AMBER, lw=2.0, arrow=True)
    rrect(slide, 320, 494, 560, 76, fill=AMBER_BG, line=AMBER, lw=1.2, adj=0.1)
    txt_cx(slide, 600, 532, '图片可诊断 且 非口腔?', 30, AMBER, bold=True, w=540)
    line(slide, 320, 532, 200, 532, INK3, lw=1.6, arrow=True)
    txt_cx(slide, 260, 522, '否', 28, INK2, w=60)
    rrect(slide, 60, 478, 200, 108, fill=WHITE, line=INK3, lw=1.0, adj=0.08)
    txt_cx(slide, 160, 516, '直接拒绝', 28, INK3, bold=True, w=190)
    txt_cx(slide, 160, 548, '共 1 次', 28, GREEN, w=190)
    txt_cx(slide, 160, 576, '质量不足', 28, INK3, w=190)
    line(slide, 600, 570, 600, 596, AMBER, lw=2.0, arrow=True)
    txt_l(slide, 612, 586, '是', 28, INK2)
    rrect(slide, 320, 596, 560, 76, fill=WHITE, line=BORDER, lw=1.0, adj=0.1)
    txt_cx(slide, 600, 632, '本地知识库检索', 32, INK, bold=True, w=520)
    txt_cx(slide, 600, 662, '无模型调用', 28, INK2, w=520)
    line(slide, 600, 672, 600, 698, AMBER, lw=2.0, arrow=True)
    rrect(slide, 320, 698, 560, 76, fill=AMBER_BG, line=AMBER, lw=1.2, adj=0.1)
    txt_cx(slide, 600, 736, '知识库强命中?', 30, AMBER, bold=True, w=540)
    line(slide, 320, 736, 200, 736, INK3, lw=1.6, arrow=True)
    txt_cx(slide, 260, 726, '是', 28, INK2, w=60)
    rrect(slide, 60, 682, 200, 108, fill=WHITE, line=GREEN, lw=1.0, adj=0.08)
    txt_cx(slide, 160, 720, '本地构建报告', 28, INK, bold=True, w=190)
    txt_cx(slide, 160, 752, '共 1 次', 28, GREEN, w=190)
    txt_cx(slide, 160, 780, '快路径', 28, INK3, w=190)
    line(slide, 600, 774, 600, 800, GREEN, lw=2.0, arrow=True)
    txt_l(slide, 612, 790, '否', 28, INK2)
    rrect(slide, 280, 800, 640, 140, fill=PANEL, line=GREEN, lw=1.4, adj=0.07)
    rect(slide, 280, 800, 4, 140, fill=GREEN)
    rrect(slide, 310, 822, 118, 40, fill=GREEN, adj=0.5)
    txt_cc(slide, 369, 843, '调用 2', 28, WHITE, bold=True, w=118)
    txt_l(slide, 448, 844, '完整诊断', 32, INK, bold=True)
    txt_l(slide, 310, 892, 'qwen3.7-plus · 主力诊断', 28, GREEN)
    txt_l(slide, 310, 924, '超时 55 秒 · 注入预检特征 · 共 2 次', 28, INK2)
    line(slide, 600, 940, 600, 968, INK3, lw=1.6, arrow=True)
    line(slide, 60, 980, 1140, 980, DIVIDER, lw=1.0, dash=True)
    txt_l(slide, 60, 1020, '云函数代理 · 前端不持有密钥', 28, INK3, bold=True)
    rrect(slide, 60, 1038, 1080, 120, fill=PANEL, line=BORDER, lw=1.0, adj=0.06)
    rect(slide, 60, 1038, 4, 120, fill=GREEN)
    txt_l(slide, 90, 1078, '云函数代理', 32, INK, bold=True)
    txt_l(slide, 90, 1114, '统一转发请求 · 密钥存环境变量 · 前端只持会话', 28, INK2)
    txt_l(slide, 90, 1146, '按阶段设超时 · 统一错误处理', 28, INK2)
    line(slide, 600, 1158, 600, 1184, INK3, lw=1.6, arrow=True)
    rrect(slide, 400, 1184, 400, 60, fill=WHITE, line=GREEN, lw=1.2, adj=0.5)
    txt_cc(slide, 600, 1214, 'DashScope 接口', 28, GREEN, bold=True, w=380)
    line(slide, 60, 1278, 1140, 1278, DIVIDER)
    txt_l(slide, 60, 1318, '模型清单', 28, INK3, bold=True)
    rrect(slide, 60, 1336, 1080, 56, fill=PANEL, adj=0.05)
    txt_cc(slide, 180, 1364, '模型', 28, INK, bold=True, w=200)
    txt_cc(slide, 430, 1364, '定位', 28, INK, bold=True, w=200)
    txt_cc(slide, 680, 1364, '用途', 28, INK, bold=True, w=240)
    txt_cc(slide, 930, 1364, '超时', 28, INK, bold=True, w=160)
    rrect(slide, 60, 1398, 1080, 68, fill=WHITE, line=BORDER, lw=1.0, adj=0.06)
    txt_cc(slide, 180, 1432, 'qwen3.6-flash', 28, GREEN, bold=True, w=220)
    txt_cc(slide, 430, 1432, '轻量预检', 28, INK2, w=200)
    txt_cc(slide, 680, 1432, '视觉预检 / 对话', 28, INK2, w=240)
    txt_cc(slide, 930, 1432, '18 秒', 28, INK2, w=160)
    rrect(slide, 60, 1474, 1080, 68, fill=WHITE, line=BORDER, lw=1.0, adj=0.06)
    txt_cc(slide, 180, 1508, 'qwen3.7-plus', 28, AMBER, bold=True, w=220)
    txt_cc(slide, 430, 1508, '主力完整诊断', 28, INK2, w=200)
    txt_cc(slide, 680, 1508, '完整诊断', 28, INK2, w=240)
    txt_cc(slide, 930, 1508, '55 秒', 28, INK2, w=160)
    line(slide, 60, 1572, 1140, 1572, DIVIDER)
    txt_l(slide, 60, 1612, '实际调用次数', 28, INK3, bold=True)
    rrect(slide, 60, 1630, 350, 130, fill=WHITE, line=GREEN, lw=1.0, adj=0.08)
    txt_l(slide, 80, 1664, '知识库强命中', 28, INK, bold=True)
    txt_l(slide, 80, 1700, '共 1 次', 28, GREEN)
    txt_l(slide, 80, 1736, '预检 + 本地报告', 28, INK2)
    rrect(slide, 425, 1630, 350, 130, fill=WHITE, line=INK3, lw=1.0, adj=0.08)
    txt_l(slide, 445, 1664, '图片不可诊断', 28, INK, bold=True)
    txt_l(slide, 445, 1700, '共 1 次', 28, INK2)
    txt_l(slide, 445, 1736, '预检 + 拒绝', 28, INK2)
    rrect(slide, 790, 1630, 350, 130, fill=WHITE, line=AMBER, lw=1.0, adj=0.08)
    txt_l(slide, 810, 1664, '完整诊断（常见）', 28, INK, bold=True)
    txt_l(slide, 810, 1700, '共 2 次', 28, AMBER)
    txt_l(slide, 810, 1736, '预检 + 主力诊断', 28, INK2)
    rrect(slide, 60, 1800, 14, 14, fill=GREEN, adj=0.2)
    txt_l(slide, 82, 1815, '模型调用', 22, INK3)
    rrect(slide, 210, 1800, 14, 14, fill=INK3, adj=0.2)
    txt_l(slide, 232, 1815, '本地处理', 22, INK3)
    rrect(slide, 350, 1800, 14, 14, fill=AMBER, adj=0.2)
    txt_l(slide, 372, 1815, '决策分支', 22, INK3)

def main():
    prs = Presentation()
    prs.slide_width = E(1200)
    prs.slide_height = E(1860)
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank); build_page1(s1)
    s2 = prs.slides.add_slide(blank); build_page2(s2)
    out = r'C:\Users\Administrator\Desktop\vetai-complete2\VetAI产品逻辑.pptx'
    prs.save(out)
    print('saved:', out)

if __name__ == '__main__':
    main()
